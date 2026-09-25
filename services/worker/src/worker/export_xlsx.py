"""Gera o XLSX do dossiê homologado a partir do snapshot (AT-021)."""
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

MONEY_FORMAT = "#,##0.00"

SHEETS = {
    "Valores": [
        ("file_id", "Arquivo"), ("ordinal", "Ordem"), ("doc_type", "Documento"), ("competence", "Competência"),
        ("section", "Seção"), ("field_key", "Campo"), ("label", "Rótulo"), ("account_code", "Conta/rubrica"),
        ("column", "Coluna"), ("value", "Valor efetivo"), ("original_value", "Valor original"),
        ("adjusted", "Ajustado"), ("nature", "D/C"), ("page", "Página"), ("bbox", "Posição (bbox)"),
    ],
    "Ajustes": [
        ("value_id", "Valor"), ("old_value", "Valor anterior"), ("new_value", "Novo valor"),
        ("reason", "Motivo"), ("author", "Autor"), ("created_at", "Data/hora"),
    ],
    "Validações": [
        ("file_id", "Arquivo"), ("rule", "Regra"), ("status", "Status"), ("expected", "Esperado"),
        ("actual", "Encontrado"), ("diff", "Diferença"), ("detail", "Detalhe"),
    ],
    "Conciliações": [
        ("competence", "Competência"), ("rule", "Regra"), ("description", "Descrição"),
        ("left_label", "Fonte A"), ("left_value", "Valor A"), ("right_label", "Fonte B"),
        ("right_value", "Valor B"), ("diff", "Diferença"), ("tolerance", "Tolerância"),
        ("status", "Status"), ("justification", "Justificativa"), ("justified_by", "Justificado por"),
        ("justified_at", "Justificado em"),
    ],
}
SOURCE_KEY = {"Valores": "values", "Ajustes": "adjustments", "Validações": "validations",
              "Conciliações": "reconciliations"}


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (list, dict)):
        return str(value)
    if hasattr(value, "tzinfo") and getattr(value, "tzinfo", None) is not None:
        return value.replace(tzinfo=None)   # openpyxl não grava datas com fuso
    return value


def build_xlsx(content: dict[str, Any], sha256: str) -> bytes:
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    case = content.get("case", {})
    company = content.get("company", {})
    for row in [
        ("Hash do snapshot (SHA-256)", sha256),
        ("Empresa", company.get("legal_name")),
        ("CNPJ", company.get("cnpj")),
        ("Período", str(case.get("period_start")) + " a " + str(case.get("period_end"))),
        ("Tolerância de conciliação (R$)", _cell(content.get("tolerance_brl"))),
        ("Homologado em", content.get("homologated_at")),
        ("Homologado por", content.get("homologated_by")),
        ("Arquivos", len(content.get("files", []))),
        ("Valores", len(content.get("values", []))),
    ]:
        resumo.append(row)
    resumo.column_dimensions["A"].width = 34
    resumo.column_dimensions["B"].width = 70

    for title, columns in SHEETS.items():
        ws = wb.create_sheet(title)
        ws.append([label for _, label in columns])
        for c in ws[1]:
            c.font = Font(bold=True)
        for item in content.get(SOURCE_KEY[title], []):
            ws.append([_cell(item.get(key)) for key, _ in columns])
        for idx, (key, _) in enumerate(columns, start=1):
            if key in ("value", "original_value", "old_value", "new_value", "expected", "actual", "diff",
                       "left_value", "right_value", "tolerance"):
                for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.number_format = MONEY_FORMAT
        ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_path(office_id: str, case_id: str, snapshot_id: str) -> str:
    return f"{office_id}/{case_id}/exports/{snapshot_id}.xlsx"


# ---------------------------------------------------------------- simulação (ciclo 2)
SIMULATION_LINE_COLUMNS = [
    ("regime", "Regime"), ("period", "Período"), ("activity", "Atividade"), ("tax", "Tributo"), ("kind", "Tipo"),
    ("base", "Base"), ("rate", "Alíquota"), ("amount", "Valor"), ("formula", "Fórmula"), ("rule_ref", "Regra@versão"),
    ("verified", "Regra conferida"), ("partial", "Trimestre parcial"), ("origin", "Origem"),
]


def build_simulation_xlsx(simulation: dict[str, Any], lines: list[dict[str, Any]], assumptions: list[dict[str, Any]]) -> bytes:
    result = simulation.get("result") or {}
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    for row in [
        ("Simulação", str(simulation.get("id"))),
        ("Hash do resultado", simulation.get("result_hash")),
        ("Hash do snapshot", simulation.get("snapshot_sha256")),
        ("Hash das premissas", simulation.get("assumptions_hash")),
        ("Regras", f"{simulation.get('rules_version')} ({simulation.get('rules_hash')})"),
        ("Competências calculadas", ", ".join(result.get("competences", []))),
        ("Competências fora do cálculo", ", ".join(result.get("excluded_competences", []))),
        ("Ordenação (elegíveis)", " < ".join(result.get("ranking", []))),
        ("Observação", "Comparativo sem recomendação: o parecer exige revisão do responsável técnico."),
    ]:
        resumo.append(row)
    resumo.column_dimensions["A"].width = 30
    resumo.column_dimensions["B"].width = 90

    comp = wb.create_sheet("Comparativo")
    comp.append(["Regime", "Status do cálculo", "Elegibilidade", "Total", "Tributos", "Pendências"])
    for c in comp[1]:
        c.font = Font(bold=True)
    elig = result.get("eligibility", {})
    for regime, r in (result.get("regimes") or {}).items():
        comp.append([regime, r.get("status"), (elig.get(regime) or {}).get("status"), _cell(Decimal(r.get("total", "0"))),
                     "; ".join(f"{k}: {v}" for k, v in (r.get("by_tax") or {}).items()), "; ".join(r.get("pending", []))])
        comp.cell(row=comp.max_row, column=4).number_format = MONEY_FORMAT

    el = wb.create_sheet("Elegibilidade")
    el.append(["Regime", "Status", "Motivo", "Regra"])
    for c in el[1]:
        c.font = Font(bold=True)
    for regime, e in elig.items():
        if not e.get("reasons"):
            el.append([regime, e.get("status"), "", ""])
        for reason in e.get("reasons", []):
            el.append([regime, reason.get("status"), reason.get("motivo"), reason.get("regra")])

    mem = wb.create_sheet("Memória")
    mem.append([label for _, label in SIMULATION_LINE_COLUMNS])
    for c in mem[1]:
        c.font = Font(bold=True)
    for line in lines:
        mem.append([_cell(line.get(key)) for key, _ in SIMULATION_LINE_COLUMNS])
        for idx in (6, 8):
            mem.cell(row=mem.max_row, column=idx).number_format = MONEY_FORMAT
        mem.cell(row=mem.max_row, column=7).number_format = "0.0000%"
    mem.freeze_panes = "A2"

    pr = wb.create_sheet("Premissas")
    pr.append(["Chave", "Escopo", "Descrição", "Sugerido", "Confirmado", "Justificativa", "Confirmado em"])
    for c in pr[1]:
        c.font = Font(bold=True)
    if not assumptions:
        pr.append(["Premissas não registradas nesta simulação (calculada antes da gravação da cópia de premissas); "
                   "recalcule para obter a memória completa"])
    for a in assumptions:
        pr.append([a.get("key"), a.get("scope"), a.get("label"), _cell(a.get("suggested_value")),
                   _cell(a.get("value")), a.get("justification"), _cell(a.get("confirmed_at"))])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def simulation_export_path(office_id: str, case_id: str, simulation_id: str) -> str:
    return f"{office_id}/{case_id}/simulations/{simulation_id}.xlsx"
