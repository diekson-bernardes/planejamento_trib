"""Integração com o Postgres local — AT-105, AT-106, AT-121, AT-122, AT-125 e exportação da memória.

Fluxo real: 6 PDFs → extração → homologação → sugestões → confirmação → cálculo → simulação → XLSX."""
import json
import time
from decimal import Decimal

import psycopg
import pytest

from conftest import SAMPLE_FILES
from worker.pipeline import Pipeline

pytestmark = [pytest.mark.db, pytest.mark.samples]
LIMIT_SECONDS = 60


def as_user(conn, user_id: str, sql: str, params: tuple = ()):
    """Executa uma RPC com auth.uid() = user_id (claims do JWT na transação)."""
    with conn.transaction():
        conn.execute("select set_config('request.jwt.claims', %s, true)",
                     (json.dumps({"sub": user_id, "role": "authenticated"}),))
        return conn.execute(sql, params).fetchall()


def prepare_case(db_conn, tenant, sample):
    with db_conn.transaction():
        db_conn.execute("insert into office_members (office_id, user_id, role) values (%s, %s, 'analyst')",
                        (tenant.office_id, tenant.user_id))
    for name in SAMPLE_FILES:
        tenant.add_file(sample(name), name + ".pdf")
    pipe = Pipeline(db_conn, tenant.storage)
    pipe.drain()
    as_user(db_conn, tenant.user_id, "select * from homologate_case(%s)", (tenant.case_id,))
    return pipe


def confirm_all(db_conn, tenant, golden, overrides: dict | None = None):
    decl = golden("motor_202608")["assumption_overrides"]
    rows = db_conn.execute("select id, key, scope, suggested_value from assumptions where case_id = %s",
                           (tenant.case_id,)).fetchall()
    for r in rows:
        value = r["suggested_value"]
        reason = None
        if r["key"] in decl:
            value, reason = decl[r["key"]], "Declaração do cliente em teste"
        if overrides and (r["key"], r["scope"]) in overrides:
            value, reason = overrides[(r["key"], r["scope"])], "Premissa alterada no teste"
        as_user(db_conn, tenant.user_id, "select confirm_assumption(%s, %s::jsonb, %s)",
                (r["id"], json.dumps(value), reason))


def simulations(db_conn, tenant):
    return db_conn.execute("select * from simulations where case_id = %s order by created_at", (tenant.case_id,)).fetchall()


def test_full_planning_flow(db_conn, tenant, sample, golden):
    pipe = prepare_case(db_conn, tenant, sample)
    as_user(db_conn, tenant.user_id, "select request_planning(%s)", (tenant.case_id,))
    pipe.drain()
    count = db_conn.execute("select count(*) as n from assumptions where case_id = %s", (tenant.case_id,)).fetchone()["n"]
    assert count >= 20

    # AT-105: pendentes bloqueiam o pedido de cálculo
    with pytest.raises(psycopg.errors.RaiseException, match="Premissas pendentes"):
        as_user(db_conn, tenant.user_id, "select request_calculation(%s)", (tenant.case_id,))

    confirm_all(db_conn, tenant, golden)
    started = time.monotonic()
    as_user(db_conn, tenant.user_id, "select request_calculation(%s)", (tenant.case_id,))
    pipe.drain()
    elapsed = time.monotonic() - started
    sims = simulations(db_conn, tenant)
    assert len(sims) == 1 and sims[0]["status"] == "done"
    assert elapsed <= LIMIT_SECONDS, f"{elapsed:.1f}s"                     # AT-125

    sim = sims[0]
    snap = db_conn.execute("select sha256 from snapshots where case_id = %s", (tenant.case_id,)).fetchone()
    assert sim["snapshot_sha256"] == snap["sha256"]                          # AT-106
    assert len(sim["assumptions_hash"]) == 64 and len(sim["rules_hash"]) == 64 and sim["rules_version"] == "2026.1.0"
    g = golden("motor_202608")
    for regime, expected in g["regimes"].items():
        assert Decimal(sim["result"]["regimes"][regime]["total"]) == Decimal(expected["total"]), regime
    assert sim["result"]["ranking"] == g["ranking"]
    lines = db_conn.execute("select count(*) as n, count(*) filter (where rule_ref = '') as sem_regra "
                            "from simulation_lines where simulation_id = %s", (sim["id"],)).fetchone()
    assert lines["n"] > 0 and lines["sem_regra"] == 0
    print(f"\n[AT-125] simulação pronta em {elapsed:.2f}s (limite {LIMIT_SECONDS}s)")

    # AT-121: mesmo pedido de novo → nenhuma simulação nova
    as_user(db_conn, tenant.user_id, "select request_calculation(%s)", (tenant.case_id,))
    pipe.drain()
    assert len(simulations(db_conn, tenant)) == 1

    # AT-122: premissa alterada → nova simulação, anterior preservada
    confirm_all(db_conn, tenant, golden, overrides={("folha.rat", "caso"): "0.03"})
    as_user(db_conn, tenant.user_id, "select request_calculation(%s)", (tenant.case_id,))
    pipe.drain()
    sims = simulations(db_conn, tenant)
    assert len(sims) == 2 and sims[0]["id"] == sim["id"] and sims[1]["result_hash"] != sim["result_hash"]

    # exportação da memória
    as_user(db_conn, tenant.user_id, "select request_simulation_export(%s)", (sim["id"],))
    pipe.drain()
    path = f"{tenant.office_id}/{tenant.case_id}/simulations/{sim['id']}.xlsx"
    assert tenant.storage.objects[path][:2] == b"PK"
    # a simulação antiga exporta as premissas com que foi calculada (RAT 0,02), não as atuais do dossiê (0,03)
    from io import BytesIO

    from openpyxl import load_workbook
    sheet = load_workbook(BytesIO(tenant.storage.objects[path]))["Premissas"]
    rat = [r for r in sheet.iter_rows(values_only=True) if r[0] == "folha.rat"]
    assert [r[4] for r in rat] == ["0.02"]

    # regenerar premissas: confirmadas com a mesma sugestão seguem confirmadas; as que deixaram de existir saem
    with db_conn.transaction():
        db_conn.execute(
            "insert into assumptions (office_id, case_id, key, scope, grp, label, value_type, suggested_value, "
            "suggested_origin, value, status) values (%s, %s, 'atividade.perfil', 'atividade:obsoleta', 'atividades', "
            "'Obsoleta', 'profile', 'null', '{}', '{}', 'confirmed')", (tenant.office_id, tenant.case_id))
    as_user(db_conn, tenant.user_id, "select request_planning(%s)", (tenant.case_id,))
    pipe.drain()
    after = db_conn.execute("select scope, status from assumptions where case_id = %s", (tenant.case_id,)).fetchall()
    assert len(after) == count and all(r["status"] == "confirmed" for r in after)
    assert not [r for r in after if r["scope"] == "atividade:obsoleta"]


def test_planning_requires_homologated_case(db_conn, tenant):
    with db_conn.transaction():
        db_conn.execute("insert into office_members (office_id, user_id, role) values (%s, %s, 'analyst')",
                        (tenant.office_id, tenant.user_id))
    with pytest.raises(psycopg.errors.RaiseException, match="Planejamento exige dossiê homologado"):
        as_user(db_conn, tenant.user_id, "select request_planning(%s)", (tenant.case_id,))
