"""AT-021: o XLSX traz todas as abas, os valores com origem e o hash do snapshot."""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from worker.export_xlsx import build_xlsx, export_path

CONTENT = {
    "schema_version": 1,
    "case": {"id": "c1", "period_start": "2026-08-01", "period_end": "2026-08-31"},
    "company": {"cnpj": "11222333000181", "legal_name": "Comércio Exemplo Ltda"},
    "tolerance_brl": Decimal("1.00"),
    "files": [{"id": "f1", "doc_type": "PGDAS_D"}],
    "values": [{"file_id": "f1", "ordinal": 0, "doc_type": "PGDAS_D", "competence": "2026-08-01",
                "section": "2.1", "field_key": "receita.rpa", "label": "Receita Bruta do PA",
                "column": "total", "value": Decimal("995.00"), "original_value": Decimal("1000.00"),
                "adjusted": True, "nature": None, "page": 1, "bbox": [1, 2, 3, 4]}],
    "adjustments": [{"value_id": "v1", "old_value": Decimal("1000.00"), "new_value": Decimal("995.00"),
                     "reason": "Correção conforme DRE", "author": "u1", "created_at": "2026-09-24T10:00:00"}],
    "validations": [{"file_id": "f1", "rule": "pgdas.soma", "status": "pass"}],
    "reconciliations": [{"competence": "2026-08-01", "rule": "R1", "status": "divergent", "diff": Decimal("10"),
                         "justification": "Devoluções lançadas em setembro"}],
    "homologated_at": "2026-09-24T10:05:00",
}
SHA = "a" * 64


def test_workbook_has_all_sheets_and_hash():
    wb = load_workbook(BytesIO(build_xlsx(CONTENT, SHA)))
    assert wb.sheetnames == ["Resumo", "Valores", "Ajustes", "Validações", "Conciliações"]
    resumo = {r[0]: r[1] for r in wb["Resumo"].iter_rows(values_only=True)}
    assert resumo["Hash do snapshot (SHA-256)"] == SHA
    assert resumo["CNPJ"] == "11222333000181"


def test_values_keep_origin_and_adjustment():
    wb = load_workbook(BytesIO(build_xlsx(CONTENT, SHA)))
    header, row = list(wb["Valores"].iter_rows(values_only=True))
    data = dict(zip(header, row))
    assert data["Valor efetivo"] == 995.0
    assert data["Valor original"] == 1000.0
    assert data["Página"] == 1
    assert data["Ajustado"] is True
    conc = dict(zip(*list(wb["Conciliações"].iter_rows(values_only=True))))
    assert conc["Justificativa"] == "Devoluções lançadas em setembro"


def test_export_path_is_inside_office_folder():
    assert export_path("o1", "c1", "s1") == "o1/c1/exports/s1.xlsx"
